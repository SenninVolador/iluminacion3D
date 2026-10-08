import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_rubric_excel():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    f_title = Font(name="Segoe UI", size=14, bold=True, color="111827")
    f_subtitle = Font(name="Segoe UI", size=10, italic=True, color="4B5563")
    f_header = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    f_subhead = Font(name="Segoe UI", size=8.5, bold=True, color="111827")
    f_data = Font(name="Segoe UI", size=9, color="1F2937")
    f_bold = Font(name="Segoe UI", size=9, bold=True, color="111827")
    f_kpi_lbl = Font(name="Segoe UI", size=8.5, bold=True, color="374151")
    f_kpi_val = Font(name="Segoe UI", size=11, bold=True, color="111827")

    fill_header_dark = PatternFill(start_color="111827", end_color="111827", fill_type="solid")
    fill_header_sub = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
    fill_header_light = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
    fill_kpi = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_nota = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    bd_thin = Side(style='thin', color='D1D5DB')
    bd_thick_bottom = Side(style='medium', color='111827')
    border_cell = Border(left=bd_thin, right=bd_thin, top=bd_thin, bottom=bd_thin)
    border_header = Border(left=bd_thin, right=bd_thin, top=bd_thin, bottom=bd_thick_bottom)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # ================= SHEET 1: Seccion_Martes =================
    ws1 = wb.create_sheet(title="Seccion_Martes")
    ws1.views.sheetView[0].showGridLines = True

    # Title & Metadata
    ws1["A1"] = "EVALUACIÓN ENTREGA N°2: ILUMINACIÓN DE INTERIOR, MOODBOARD Y RÉPLICA — SECCIÓN MARTES"
    ws1["A1"].font = f_title
    ws1["A2"] = "Asignatura: Iluminación 3D y Shaders · Docente: Daniel Rojas (UNIACC) · Escala: 1.0 a 7.0 (Exigencia 60%, 36 pts = 4.0)"
    ws1["A2"].font = f_subtitle

    # Course Stats / KPI Cards
    ws1["B4"] = "Promedio Curso:"
    ws1["B4"].font = f_kpi_lbl
    ws1["B4"].alignment = align_right
    ws1["C4"] = "=AVERAGE(I6:I18)"
    ws1["C4"].font = f_kpi_val
    ws1["C4"].alignment = align_center
    ws1["C4"].fill = fill_kpi
    ws1["C4"].border = border_cell
    ws1["C4"].number_format = "0.0"

    ws1["D4"] = "Tasa Aprobación:"
    ws1["D4"].font = f_kpi_lbl
    ws1["D4"].alignment = align_right
    ws1["E4"] = '=COUNTIF(J6:J18,"APROBADO")/COUNTA(B6:B18)'
    ws1["E4"].font = f_kpi_val
    ws1["E4"].alignment = align_center
    ws1["E4"].fill = fill_kpi
    ws1["E4"].border = border_cell
    ws1["E4"].number_format = "0.0%"

    ws1["F4"] = "Nota Máxima:"
    ws1["F4"].font = f_kpi_lbl
    ws1["F4"].alignment = align_right
    ws1["G4"] = "=MAX(I6:I18)"
    ws1["G4"].font = f_kpi_val
    ws1["G4"].alignment = align_center
    ws1["G4"].fill = fill_kpi
    ws1["G4"].border = border_cell
    ws1["G4"].number_format = "0.0"

    ws1["H4"] = "Nota Mínima:"
    ws1["H4"].font = f_kpi_lbl
    ws1["H4"].alignment = align_right
    ws1["I4"] = "=MIN(I6:I18)"
    ws1["I4"].font = f_kpi_val
    ws1["I4"].alignment = align_center
    ws1["I4"].fill = fill_kpi
    ws1["I4"].border = border_cell
    ws1["I4"].number_format = "0.0"

    # Table Headers (Row 5)
    headers = [
        ("N°", 5),
        ("Estudiante (Sección Martes)", 36),
        ("1. Moodboard\ny Referencias\n(Máx 9 pts / 15%)", 17),
        ("2. Iluminación\nInterior\n(Máx 15 pts / 25%)", 17),
        ("3. Uso Correcto\nde Luces\n(Máx 18 pts / 30%)", 17),
        ("4. Intención\ny Atmósfera\n(Máx 12 pts / 20%)", 17),
        ("5. Nomenclatura\ny Orden\n(Máx 6 pts / 10%)", 16),
        ("Puntaje\nTotal\n(Máx 60)", 11),
        ("Nota Final\n(Escala 1.0 - 7.0)", 13),
        ("Estado\nAcadémico", 13),
        ("Observaciones y Retroalimentación Técnica", 44)
    ]

    for col_idx, (h_text, width) in enumerate(headers, start=1):
        cell = ws1.cell(row=5, column=col_idx, value=h_text)
        cell.font = f_header
        cell.fill = fill_header_dark
        cell.alignment = align_center
        cell.border = border_header
        ws1.column_dimensions[get_column_letter(col_idx)].width = width

    ws1.row_dimensions[5].height = 42

    # Students List (The 13 students)
    students = [
        "JOSE MIGUEL FARFAN CARRASCO",
        "BENJAMIN FERNANDO FIGUEROA GONZALEZ",
        "MATIAS IGNACIO GODOY TIRADO",
        "HENRY ALEXANDER HERNANDEZ PEREZ",
        "GAEL ALBERTO IBARRA MONTECINOS",
        "LUKA PASCAL LATIFE CUEVAS",
        "MEI MALDONADO DUARTE",
        "ELIZABETH FABIOLA MEDRANO MORAN",
        "FREDDY DANILO MONTANARES ROSALES",
        "JOSE FRANCISCO NAVARRETE BARRIOS",
        "NATANIEL ASTRO REYES FLORES",
        "DANIEL ELIAS ROJAS LOPEZ",
        "CRISTOBAL NICOLAS VALENZUELA PRESSAC"
    ]

    for idx, name in enumerate(students, start=1):
        r = 5 + idx
        ws1.row_dimensions[r].height = 24
        zebra = (idx % 2 == 0)

        # Col A: N°
        c_num = ws1.cell(row=r, column=1, value=idx)
        c_num.font = f_bold
        c_num.alignment = align_center
        c_num.border = border_cell
        if zebra: c_num.fill = fill_zebra

        # Col B: Estudiante
        c_name = ws1.cell(row=r, column=2, value=name)
        c_name.font = f_data
        c_name.alignment = align_left
        c_name.border = border_cell
        if zebra: c_name.fill = fill_zebra

        # Cols C to G: Evaluation Criteria Scores (initially blank for teacher input)
        for c_idx in range(3, 8):
            cell = ws1.cell(row=r, column=c_idx)
            cell.font = f_data
            cell.alignment = align_center
            cell.border = border_cell
            if zebra: cell.fill = fill_zebra

        # Col H: Puntaje Total
        c_tot = ws1.cell(row=r, column=8, value=f"=SUM(C{r}:G{r})")
        c_tot.font = f_bold
        c_tot.alignment = align_center
        c_tot.border = border_cell
        if zebra: c_tot.fill = fill_zebra

        # Col I: Nota (Official Chilean formula at 60% requirement: 36 pts = 4.0)
        c_nota = ws1.cell(row=r, column=9, value=f'=IF(H{r}="","",ROUND(IF(H{r}<36, 1 + 3*(H{r}/36), 4 + 3*((H{r}-36)/(60-36))), 1))')
        c_nota.font = f_bold
        c_nota.alignment = align_center
        c_nota.fill = fill_nota
        c_nota.border = border_cell
        c_nota.number_format = "0.0"

        # Col J: Estado
        c_est = ws1.cell(row=r, column=10, value=f'=IF(I{r}="","",IF(I{r}>=4.0,"APROBADO","REPROBADO"))')
        c_est.font = f_bold
        c_est.alignment = align_center
        c_est.border = border_cell
        if zebra: c_est.fill = fill_zebra

        # Col K: Observaciones
        c_obs = ws1.cell(row=r, column=11)
        c_obs.font = f_data
        c_obs.alignment = align_left
        c_obs.border = border_cell
        if zebra: c_obs.fill = fill_zebra


    # ================= SHEET 2: Matriz_Rubrica =================
    ws2 = wb.create_sheet(title="Matriz_Rubrica")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = "MATRIZ ANALÍTICA DE EVALUACIÓN: ENTREGA N°2 — ILUMINACIÓN DE INTERIOR"
    ws2["A1"].font = f_title
    ws2["A2"] = "Desglose por niveles de desempeño para cada criterio de evaluación (Total: 60 Puntos = 100%)"
    ws2["A2"].font = f_subtitle

    matriz_headers = [
        ("Criterio y Ponderación", 25),
        ("Excelente (100%)", 30),
        ("Bueno (75%)", 30),
        ("Suficiente (50%)", 30),
        ("Insuficiente (0-25%)", 30),
        ("Pts Máx", 10),
        ("Pond (%)", 10)
    ]

    for col_idx, (h_text, width) in enumerate(matriz_headers, start=1):
        cell = ws2.cell(row=4, column=col_idx, value=h_text)
        cell.font = f_header
        cell.fill = fill_header_dark
        cell.alignment = align_center
        cell.border = border_header
        ws2.column_dimensions[get_column_letter(col_idx)].width = width

    ws2.row_dimensions[4].height = 28

    rubric_rows = [
        (
            "1. Referencias Visuales, Análisis y Moodboard",
            "Moodboard estructurado con referencias de alta calidad cinematográfica/videojuegos AAA. Desglose analítico riguroso (fuentes, temperatura Kelvin, clave tonal). La escena 3D refleja fielmente la dirección visual propuesta.",
            "Moodboard completo pero con desglose técnico elemental o leve discrepancia en la traslación de la paleta y contrastes a la escena 3D.",
            "Moodboard escaso o desorganizado; recopilación de imágenes aisladas sin análisis técnico de fuentes ni intención clara.",
            "No presenta moodboard o incluye imágenes no pertinentes sin relación con la escena desarrollada.",
            9,
            "15%"
        ),
        (
            "2. Iluminación Interior y Espacialidad",
            "Recinto con excelente lectura volumétrica y planos de profundidad (primer plano, medio y fondo). Luz motivada por aberturas. Rebotes indirectos creíbles (Lumen) sin empastar sombras ni lavar negros con luz falsa.",
            "Espacio bien iluminado con sensación de volumen, aunque con leve aplanamiento en rincones o rebotes indirectos ligeramente desbalanceados.",
            "Lectura espacial deficiente; zonas excesivamente oscuras sin detalle o iluminación plana que debilita la tridimensionalidad del interior.",
            "Espacio interior sin profundidad, sobreexpuesto en su totalidad o en penumbra ilegible por falta de rebotes.",
            15,
            "25%"
        ),
        (
            "3. Uso Correcto y Técnico de Luces",
            "Selección idónea de emisores (Rect Lights en vanos, Point/Spot con Source Radius en bombillas). Jerarquía clara según Brejon. Atenuación física realista, sombras calibradas y balance térmico Kelvin impecable.",
            "Tipos de luces adecuados y bien ubicados, pero con desajustes menores en radio de atenuación, sombras algo duras o leve desbalance Kelvin.",
            "Uso incorrecto de tipos de luz (ej. Point Lights en ventanas), sombras con artefactos o intensidades arbitrarias sin jerarquía física.",
            "Luces incoherentes, fugas de luz (light leaks) a través de muros o esquema plano sin emisores justificados.",
            18,
            "30%"
        ),
        (
            "4. Intención Narrativa y Atmósfera",
            "Atmósfera dramática definida y envolvente (tensión, calidez, melancolía, etc.). Punto focal nítido donde la luz guía la mirada de forma orgánica hacia el centro de interés. Niebla volumétrica usada con sentido estético.",
            "Atmósfera presente y perceptible, pero el punto focal compite con otros brillos secundarios o la niebla volumétrica satura ligeramente el encuadre.",
            "Intención confusa; la iluminación no transmite un tono emocional claro ni logra conducir la atención hacia un elemento protagónico.",
            "Escena sin atmósfera ni intención narrativa; iluminación puramente utilitaria sin propósito estético ni jerarquía visual.",
            12,
            "20%"
        ),
        (
            "5. Nomenclatura, Orden y Entregables",
            "Nomenclatura oficial UE5 respetada en la totalidad de actores. Outliner limpio organizado en carpetas temáticas. Post-Process calibrado sin quemar blancos. Cuatro capturas HD y lámina entregadas impecablemente.",
            "Proyecto ordenado con descuidos menores en nombres de luces o carpetas. Post-Process correcto y capturas completas.",
            "Desorden notorio en Outliner (nombres por defecto PointLight_1), falta una captura solicitada o exposición inestable.",
            "Proyecto desorganizado, archivos faltantes, sin carpetas ni nomenclatura estándar, o faltan entregables clave.",
            6,
            "10%"
        )
    ]

    for idx, (crit, ex, bu, suf, ins, pts, pond) in enumerate(rubric_rows, start=5):
        ws2.row_dimensions[idx].height = 65
        zebra = (idx % 2 == 0)

        c1 = ws2.cell(row=idx, column=1, value=crit)
        c1.font = f_bold
        c1.alignment = align_left
        c1.border = border_cell

        for col_i, text in enumerate([ex, bu, suf, ins], start=2):
            c = ws2.cell(row=idx, column=col_i, value=text)
            c.font = f_data
            c.alignment = align_left
            c.border = border_cell
            if zebra: c.fill = fill_zebra

        c_p = ws2.cell(row=idx, column=6, value=pts)
        c_p.font = f_bold
        c_p.alignment = align_center
        c_p.border = border_cell
        c_p.fill = fill_header_light

        c_po = ws2.cell(row=idx, column=7, value=pond)
        c_po.font = f_bold
        c_po.alignment = align_center
        c_po.border = border_cell
        c_po.fill = fill_header_light


    # ================= SHEET 3: Escala_Conversion =================
    ws3 = wb.create_sheet(title="Escala_Conversion")
    ws3.views.sheetView[0].showGridLines = True

    ws3["A1"] = "TABLA OFICIAL DE CONVERSIÓN: PUNTAJE A NOTA (ESCALA CHILENA 1.0 - 7.0 AL 60%)"
    ws3["A1"].font = f_title
    ws3["A2"] = "Puntaje Máximo: 60 Puntos · Puntaje de Aprobación (Nota 4.0): 36 Puntos"
    ws3["A2"].font = f_subtitle

    scale_headers = [
        ("Puntaje", 12),
        ("Nota Calculada", 16),
        ("Estado Académico", 18),
        ("Nivel de Desempeño y Criterio", 42)
    ]

    for col_idx, (h_text, width) in enumerate(scale_headers, start=1):
        cell = ws3.cell(row=4, column=col_idx, value=h_text)
        cell.font = f_header
        cell.fill = fill_header_dark
        cell.alignment = align_center
        cell.border = border_header
        ws3.column_dimensions[get_column_letter(col_idx)].width = width

    ws3.row_dimensions[4].height = 26

    # Generate row for each integer point from 60 down to 0
    for p in range(60, -1, -1):
        r = 4 + (60 - p) + 1
        ws3.row_dimensions[r].height = 19
        zebra = ((60 - p) % 2 == 0)

        # Score
        c_p = ws3.cell(row=r, column=1, value=p)
        c_p.font = f_bold
        c_p.alignment = align_center
        c_p.border = border_cell
        if zebra: c_p.fill = fill_zebra

        # Grade
        if p >= 36:
            grade = round(4.0 + 3.0 * ((p - 36) / (60 - 36)), 1)
            status = "APROBADO"
        else:
            grade = round(1.0 + 3.0 * (p / 36), 1)
            status = "REPROBADO"

        c_g = ws3.cell(row=r, column=2, value=grade)
        c_g.font = f_bold
        c_g.alignment = align_center
        c_g.border = border_cell
        c_g.fill = fill_nota if grade >= 4.0 else fill_zebra
        c_g.number_format = "0.0"

        c_s = ws3.cell(row=r, column=3, value=status)
        c_s.font = f_bold
        c_s.alignment = align_center
        c_s.border = border_cell
        if zebra: c_s.fill = fill_zebra

        # Descriptor
        if p == 60:
            desc = "Sobresaliente: Dominio técnico impecable y réplica cinematográfica profesional."
        elif p >= 54:
            desc = "Muy Bueno: Excelente nivel de iluminación y moodboard sólido."
        elif p >= 48:
            desc = "Bueno: Cumplimiento consistente de todos los requerimientos."
        elif p >= 42:
            desc = "Aceptable: Cumple los requisitos con detalles de calibración."
        elif p >= 36:
            desc = "Aprobado (Corte 60%): Cumplimiento mínimo de los objetivos formativos."
        elif p >= 24:
            desc = "Reprobado: Falencias graves en jerarquía lumínica o ausencia de moodboard."
        else:
            desc = "Insuficiente: Incumplimiento crítico de los requisitos de entrega."

        c_d = ws3.cell(row=r, column=4, value=desc)
        c_d.font = f_data
        c_d.alignment = align_left
        c_d.border = border_cell
        if zebra: c_d.fill = fill_zebra

    dest1 = r"c:\Users\danie\OneDrive\Escritorio\Rubrica_Evaluacion_Entrega_02.xlsx"
    dest2 = r"c:\Users\danie\OneDrive\Escritorio\Iluminacion3D\docs\public\Rubrica_Evaluacion_Entrega_02.xlsx"
    wb.save(dest1)
    wb.save(dest2)
    print(f"Workbook saved to {dest1} and {dest2}")

if __name__ == "__main__":
    create_rubric_excel()
